# -*- coding: utf-8 -*-

"""
    项目接口文档写回：将已确认的枚举值覆盖写入上传的Excel。
"""

from __future__ import annotations

import os
import tempfile
from copy import copy
from dataclasses import dataclass
from typing import Any, Dict, Mapping, Optional
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.cell_range import CellRange
from openpyxl.worksheet.worksheet import Worksheet

from applications.data_generation.services.autotest_interface_document_parser import (
    MAX_DOCUMENT_COLUMNS,
    MAX_DOCUMENT_SIZE,
)
from configure import PROJECT_CONFIG

_LENGTH_HEADERS = {"字段长度", "长度"}
_REMARK_HEADERS = {"备注"}
_ENUM_HEADER = "枚举值"
_ENUM_ROW_HEIGHT_PER_LINE = 18


class ProjectInterfaceWorkbookError(ValueError):
    """项目接口文档无法安全改写。"""


@dataclass(frozen=True)
class ProjectInterfaceWorkbookOverwriteResult:
    """
        封装改写完成后的结果
    """
    file_path: str          #原始文件路径
    enum_column: int        #枚举值所在列
    inserted_column: bool   #是否插入枚举值列
    written_count: int      #实际写入枚举值的字段数量
    file_size: int          #覆盖后的文件大小


def resolve_uploaded_interface_document_path(storage_key: str) -> str:
    """
        把数据库任务表中保存的存储键转换为本地绝对路径
    """
    key = str(storage_key or "").strip().replace("\\", "/").lstrip("/")
    if not key or ".." in key.split("/"):
        raise ProjectInterfaceWorkbookError("接口文档存储路径不合法")
    root = os.path.abspath(PROJECT_CONFIG.OUTPUT_UPLOAD_DIR)
    target = os.path.abspath(os.path.join(root, *key.split("/")))
    real_root = os.path.realpath(root)
    real_target = os.path.realpath(target)
    try:
        is_inside_root = os.path.commonpath((real_root, real_target)) == real_root
    except ValueError as exc:
        raise ProjectInterfaceWorkbookError("接口文档存储路径越界") from exc
    if target == root or not is_inside_root:
        raise ProjectInterfaceWorkbookError("接口文档存储路径越界")
    return target


def _header_text(value: object) -> str:
    """ 统一处理表头文字，去掉表头中的空白字符 """
    return "".join(str(value or "").replace("\xa0", " ").split())


def _find_column(sheet: Worksheet, header_row: int, candidates: set[str]) -> int:
    """
        在表头查找制定名称的列
    """
    columns = [
        column
        for column in range(1, sheet.max_column + 1)
        if _header_text(sheet.cell(header_row, column).value) in candidates
    ]
    if len(columns) != 1:
        names = "/".join(sorted(candidates))
        raise ProjectInterfaceWorkbookError(f"项目接口文档表头[{names}]不唯一或不存在")
    return columns[0]


def _capture_column_dimensions(sheet: Worksheet, start_column: int) -> Dict[int, Dict[str, Any]]:
    """
        保存枚举插入位置的显示属性
    """
    captured: Dict[int, Dict[str, Any]] = {}
    for column in range(start_column, sheet.max_column + 1):
        dimension = sheet.column_dimensions.get(get_column_letter(column))
        if dimension is None:
            continue
        captured[column] = {
            "width": dimension.width,
            "hidden": dimension.hidden,
            "bestFit": dimension.bestFit,
            "outlineLevel": dimension.outlineLevel,
            "collapsed": dimension.collapsed,
        }
    return captured


def _restore_shifted_column_dimensions(
        sheet: Worksheet,
        captured: Mapping[int, Mapping[str, Any]],
) -> None:
    """
        恢复移动后的列属性
    """
    for source_column in sorted(captured, reverse=True):
        target = sheet.column_dimensions[get_column_letter(source_column + 1)]
        for name, value in captured[source_column].items():
            setattr(target, name, value)


def _shift_merged_ranges(sheet: Worksheet, insert_column: int) -> list[CellRange]:
    """
        处理合并单元格
    """
    shifted: list[CellRange] = []
    for merged_range in list(sheet.merged_cells.ranges):
        cell_range = CellRange(str(merged_range))
        sheet.unmerge_cells(str(merged_range))
        if cell_range.min_col >= insert_column:
            cell_range.shift(col_shift=1)
        elif cell_range.max_col >= insert_column:
            cell_range.max_col += 1
        shifted.append(cell_range)
    return shifted


def _insert_enum_column(sheet: Worksheet, insert_column: int) -> None:
    """
        将枚举值插入到新列
    """
    dimensions = _capture_column_dimensions(sheet, insert_column)
    merged_ranges = _shift_merged_ranges(sheet, insert_column)
    sheet.insert_cols(insert_column, 1)
    _restore_shifted_column_dimensions(sheet, dimensions)
    sheet.column_dimensions[get_column_letter(insert_column)].width = 24

    # 枚举值和备注同为文本内容，复制插入后右侧备注列的逐行样式。
    style_source_column = insert_column + 1
    for row in range(1, sheet.max_row + 1):
        sheet.cell(row, insert_column)._style = copy(
            sheet.cell(row, style_source_column)._style
        )
    for cell_range in merged_ranges:
        sheet.merge_cells(str(cell_range))


def _verify_written_workbook(
        file_path: str,
        *,
        header_row: int,
        enum_column: int,
        expected_values: Mapping[int, Optional[str]],
) -> None:
    """
        重新打开临时保存的excel表，验证是否成功写入枚举值
    """
    workbook = load_workbook(file_path, read_only=True, data_only=False)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        if _header_text(sheet.cell(header_row, enum_column).value) != _ENUM_HEADER:
            raise ProjectInterfaceWorkbookError("枚举值列写入校验失败")
        for row, expected in expected_values.items():
            actual = sheet.cell(row, enum_column).value
            if (None if actual is None else str(actual)) != expected:
                raise ProjectInterfaceWorkbookError(f"第{row}行枚举值写入校验失败")
    finally:
        workbook.close()


class AutoTestProjectInterfaceWorkbookService:
    """
        将枚举抽取结果写回任务本地项目文档并覆盖原文件。
    """

    def overwrite_project_document(
            self,
            storage_key: str,
            interface_document: Mapping[str, Any],
    ) -> ProjectInterfaceWorkbookOverwriteResult:
        """
            外部真正调用功能的方法
            storage_key：原始接口文档在上传目录汇总的相对路径
            interface_document：枚举提取后的项目接口文档快照
        """
        if str(interface_document.get("interface_style") or "").strip().lower() != "project":
            raise ProjectInterfaceWorkbookError("只允许改写项目接口文档")
        fields = interface_document.get("fields")
        header_row = interface_document.get("header_row")
        if not isinstance(fields, list) or not isinstance(header_row, int) or header_row < 1:
            raise ProjectInterfaceWorkbookError("项目接口文档快照结构不正确")

        file_path = resolve_uploaded_interface_document_path(storage_key)
        if not file_path.lower().endswith(".xlsx"):
            raise ProjectInterfaceWorkbookError("项目接口文档仅支持xlsx格式")
        if os.path.islink(file_path) or not os.path.isfile(file_path):
            raise ProjectInterfaceWorkbookError("项目接口文档不存在或不是普通文件")
        original_stat = os.stat(file_path)
        if original_stat.st_size < 1 or original_stat.st_size > MAX_DOCUMENT_SIZE:
            raise ProjectInterfaceWorkbookError("项目接口文档大小超出限制")

        temp_path: Optional[str] = None
        workbook = None
        try:
            workbook = load_workbook(file_path, read_only=False, data_only=False)
            if not workbook.sheetnames:
                raise ProjectInterfaceWorkbookError("项目接口文档不包含工作表")
            sheet = workbook[workbook.sheetnames[0]]
            if sheet.max_column > MAX_DOCUMENT_COLUMNS:
                raise ProjectInterfaceWorkbookError("项目接口文档列数超出限制")
            if header_row > sheet.max_row:
                raise ProjectInterfaceWorkbookError("项目接口文档表头行越界")

            length_column = _find_column(sheet, header_row, _LENGTH_HEADERS)
            remark_column = _find_column(sheet, header_row, _REMARK_HEADERS)
            enum_columns = [
                column
                for column in range(1, sheet.max_column + 1)
                if _header_text(sheet.cell(header_row, column).value) == _ENUM_HEADER
            ]
            if len(enum_columns) > 1:
                raise ProjectInterfaceWorkbookError("项目接口文档存在多个枚举值列")

            inserted_column = not enum_columns
            if inserted_column:
                if sheet.max_column >= MAX_DOCUMENT_COLUMNS:
                    raise ProjectInterfaceWorkbookError("项目接口文档已达最大列数，无法添加枚举值列")
                if length_column >= remark_column:
                    raise ProjectInterfaceWorkbookError("字段长度列必须位于备注列之前")
                enum_column = remark_column
                _insert_enum_column(sheet, enum_column)
            else:
                enum_column = enum_columns[0]
                if not length_column < enum_column < remark_column:
                    raise ProjectInterfaceWorkbookError("枚举值列必须位于字段长度与备注之间")

            sheet.cell(header_row, enum_column).value = _ENUM_HEADER
            expected_values: Dict[int, Optional[str]] = {}
            written_count = 0
            for field in fields:
                if not isinstance(field, Mapping):
                    raise ProjectInterfaceWorkbookError("项目接口字段快照结构不正确")
                source_row = field.get("source_row")
                if (
                        not isinstance(source_row, int)
                        or source_row <= header_row
                        or source_row > sheet.max_row
                ):
                    raise ProjectInterfaceWorkbookError("项目接口字段行号越界")
                raw_value = field.get("enum_values")
                value = str(raw_value).strip() if raw_value is not None else None
                value = value or None
                cell = sheet.cell(source_row, enum_column)
                cell.value = value
                if value is not None:
                    alignment = copy(cell.alignment)
                    alignment.wrap_text = True
                    cell.alignment = alignment
                    required_height = (value.count("\n") + 1) * _ENUM_ROW_HEIGHT_PER_LINE
                    row_dimension = sheet.row_dimensions[source_row]
                    if required_height > (row_dimension.height or 0):
                        row_dimension.height = required_height
                expected_values[source_row] = value
                written_count += value is not None

            file_descriptor, temp_path = tempfile.mkstemp(
                prefix=".project-enum-",
                suffix=".xlsx",
                dir=os.path.dirname(file_path),
            )
            os.close(file_descriptor)
            workbook.save(temp_path)
            workbook.close()
            workbook = None
            _verify_written_workbook(
                temp_path,
                header_row=header_row,
                enum_column=enum_column,
                expected_values=expected_values,
            )
            os.chmod(temp_path, original_stat.st_mode & 0o777)
            os.replace(temp_path, file_path)
            temp_path = None
            return ProjectInterfaceWorkbookOverwriteResult(
                file_path=file_path,
                enum_column=enum_column,
                inserted_column=inserted_column,
                written_count=written_count,
                file_size=os.path.getsize(file_path),
            )
        except ProjectInterfaceWorkbookError:
            raise
        except (BadZipFile, InvalidFileException, OSError, ValueError, KeyError) as exc:
            raise ProjectInterfaceWorkbookError("项目接口文档改写失败") from exc
        finally:
            if workbook is not None:
                workbook.close()
            if temp_path and os.path.isfile(temp_path):
                try:
                    os.remove(temp_path)
                except OSError:
                    pass
