# -*- coding: utf-8 -*-

'''
    读取xlsx类型的接口文档的第一个sheet，并把表格中的字段转换为统一的python字典结构
'''

from __future__ import annotations

import io
import re
from dataclasses import dataclass
from typing import Callable, Dict, Iterator, List, Optional, Tuple, TypedDict
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.worksheet import Worksheet

from applications.data_generation.constants import (
    MAX_DOCUMENT_COLUMNS,
    MAX_DOCUMENT_ROWS,
    MAX_DOCUMENT_SIZE,
)
from applications.data_generation.services.autotest_interface_document_headers import (
    ENUM_COLUMN_HEADER,
    ESB_HEADERS,
    INTEGRATION_HEADERS,
    PROJECT_HEADERS,
    cell_text,
    enum_header_matches,
    esb_header_matches,
    integration_collection_header_level,
    integration_header_matches,
    project_header_matches,
)
from enums import AutoTestInterfaceStyle

#必填值和非必填
_REQUIRED_TRUE_VALUES = {"1", "true", "m", "y", "yes", "是", "必填", "必输"}
_REQUIRED_FALSE_VALUES = {"0", "false", "o", "n", "no", "否", "非必填", "非必输"}
#数组类型
_ARRAY_TYPE_VALUES = {"array", "list"}
_STRUCT_TYPE_VALUE = "struct"
_FIELD_NAME_PATTERN = re.compile(r"[A-Za-z0-9]+")


class InterfaceDocumentParseError(ValueError):
    """接口文档结构或内容不符合约定。"""


class ParsedInterfaceField(TypedDict):
    """保存接口文档中的叶子字段及其规则和父级容器路径。

    解析器生成该结构，字段匹配与场景生成读取它；array_path 保留旧快照键名，
    现在按顺序包含数组与结构体节点。
    """
    field_name: str
    field_chinese_name: Optional[str]
    required: Optional[bool]
    required_text: Optional[str]
    length: Optional[str]
    total_length: Optional[int]
    integer_length: Optional[int]
    decimal_length: Optional[int]
    field_name_error: Optional[str]
    length_error: Optional[str]
    enum_values: Optional[str]          #枚举文本：esb直接读取，其他文档由AI抽取
    remark: Optional[str]               #AI枚举抽取使用的原始字段文本
    array_path: List[str]               #兼容旧快照：数组与结构体的完整父级路径
    duplicate_in_context: bool          #字段是否在同一个上下文中重复
    source_row: int                     #字段在原excel中的行号


class ParsedInterfaceDocument(TypedDict):
    """
        整份文档输出结构
    """
    interface_style: str
    sheet_name: str
    header_row: int
    fields: List[ParsedInterfaceField]


class _ContainerContext(TypedDict):
    """记录当前未配对的数组或结构体标记，供标准文档按嵌套顺序配对。"""

    name: str
    kind: str
    start_row: int


_FieldIdentity = Tuple[Tuple[str, ...], str]


@dataclass(frozen=True)
class _StandardSheetConfig:
    interface_style: AutoTestInterfaceStyle
    headers: Tuple[str, ...]
    matches_header: Callable[[object, str], bool]
    field_name_header: str
    chinese_name_header: str
    array_type_header: str
    required_header: str
    length_header: str
    enum_header: Optional[str] = None
    remark_header: Optional[str] = None


_STANDARD_SHEET_CONFIGS = {
    AutoTestInterfaceStyle.ESB: _StandardSheetConfig(
        interface_style=AutoTestInterfaceStyle.ESB,
        headers=ESB_HEADERS,
        matches_header=esb_header_matches,
        field_name_header="英文名称",
        chinese_name_header="中文名称",
        array_type_header="数据格式",
        required_header="是否必输",
        length_header="长度",
        enum_header=ENUM_COLUMN_HEADER,
    ),
    AutoTestInterfaceStyle.PROJECT: _StandardSheetConfig(
        interface_style=AutoTestInterfaceStyle.PROJECT,
        headers=PROJECT_HEADERS,
        matches_header=project_header_matches,
        field_name_header="字段名",
        chinese_name_header="数标",
        array_type_header="格式",
        required_header="必填",
        length_header="字段长度",
        remark_header="备注",
    ),
}


def _normalize_required(value: Optional[str]) -> Optional[bool]:
    if value is None:
        return None
    normalized = re.sub(r"\s+", "", value.casefold())
    if normalized in _REQUIRED_TRUE_VALUES:
        return True
    if normalized in _REQUIRED_FALSE_VALUES:
        return False
    return None


def _normalize_length(
        value: Optional[str],
) -> Tuple[Optional[str], Optional[int], Optional[int], Optional[int]]:
    if value is None:
        return None, None, None, None
    normalized = re.sub(r"\s*[,，]\s*", ",", value.strip())
    decimal_match = re.fullmatch(r"(\d+),(\d+)", normalized)
    if decimal_match:
        total_length = int(decimal_match.group(1))
        decimal_length = int(decimal_match.group(2))
        return normalized, total_length, total_length - decimal_length, decimal_length
    integer_match = re.fullmatch(r"\d+", normalized)
    if integer_match:
        total_length = int(normalized)
        return normalized, total_length, total_length, None
    return normalized, None, None, None


def _validate_workbook_columns(sheet: Worksheet) -> None:
    if sheet.max_column > MAX_DOCUMENT_COLUMNS:
        raise InterfaceDocumentParseError(f"接口文档列数不能超过{MAX_DOCUMENT_COLUMNS}列")


def _iter_limited_rows(
        sheet: Worksheet,
        *,
        min_row: int = 1,
) -> Iterator[Tuple[int, Tuple[object, ...], bool]]:
    """遍历工作表并按实际有值的行数执行资源限制。"""
    content_row_count = 0
    for row_no, row in enumerate(
            sheet.iter_rows(min_row=min_row, values_only=True),
            start=min_row,
    ):
        is_empty = not any(cell_text(value) is not None for value in row)
        if not is_empty:
            content_row_count += 1
            if content_row_count > MAX_DOCUMENT_ROWS:
                raise InterfaceDocumentParseError(
                    f"接口文档行数不能超过{MAX_DOCUMENT_ROWS}行"
                )
        yield row_no, row, is_empty


def _iter_content_rows(
        sheet: Worksheet,
        *,
        min_row: int = 1,
) -> Iterator[Tuple[int, Tuple[object, ...]]]:
    """跳过空行并返回有效内容行。"""
    for row_no, row, is_empty in _iter_limited_rows(sheet, min_row=min_row):
        if is_empty:
            continue
        yield row_no, row


_HeaderMatcher = Callable[[object, str], bool]


def _find_ordered_header_group(
        sheet: Worksheet,
        expected_headers: Tuple[str, ...],
        matches_header: _HeaderMatcher,
        *,
        document_name: str = "接口文档",
) -> Tuple[int, Tuple[object, ...], Dict[str, int]]:
    """
        定位首个顺序一致的表头组；目标列之间允许夹杂无关列。
    """
    anchor_seen = False
    for row_no, row in _iter_content_rows(sheet):
        anchor_seen = anchor_seen or any(
            matches_header(value, expected_headers[0]) for value in row
        )
        positions: Dict[str, int] = {}
        next_column = 0
        for expected_header in expected_headers:
            column_no = next(
                (
                    index
                    for index in range(next_column, len(row))
                    if matches_header(row[index], expected_header)
                ),
                None,
            )
            if column_no is None:
                break
            positions[expected_header] = column_no
            next_column = column_no + 1
        if len(positions) == len(expected_headers):
            return row_no, row, positions

    if anchor_seen:
        raise InterfaceDocumentParseError(
            f"{document_name}字段行须按顺序包含: {', '.join(expected_headers)}"
        )
    raise InterfaceDocumentParseError(
        f"未找到{document_name}字段行标识「{expected_headers[0]}」"
    )


def _row_values(row: Tuple[object, ...], positions: Dict[str, int]) -> Dict[str, Optional[str]]:
    return {
        header: cell_text(row[column_no]) if column_no < len(row) else None
        for header, column_no in positions.items()
    }


def _iter_field_rows(
        sheet: Worksheet,
        *,
        header_row: int,
        positions: Dict[str, int],
        field_name_header: str,
        chinese_name_header: str,
) -> Iterator[Tuple[int, Dict[str, Optional[str]]]]:
    """从首个有效字段开始读取，遇到空行或非英数字字段名时停止。

    :param sheet: 待扫描的接口文档工作表。
    :param header_row: 字段表头所在的行号。
    :param positions: 表头名称到工作表列索引的映射。
    :param field_name_header: 英文字段名对应的表头名称。
    :param chinese_name_header: 中文字段名对应的表头名称，用于确认扫描起点。
    :return: 按工作表顺序生成字段行号及其各列文本值。
    """
    data_started = False
    for row_no, row, is_empty in _iter_limited_rows(
            sheet,
            min_row=header_row + 1,
    ):
        if is_empty:
            if data_started:
                break
            continue

        values = _row_values(row, positions)
        field_name = values[field_name_header]
        is_alphanumeric_field_name = _FIELD_NAME_PATTERN.fullmatch(field_name or "") is not None
        if data_started:
            if not is_alphanumeric_field_name:
                break
        else:
            chinese_name = values[chinese_name_header]
            if not (is_alphanumeric_field_name and chinese_name):
                continue
            data_started = True

        yield row_no, values


def _update_container_contexts(
        active_containers: List[_ContainerContext],
        *,
        container_name: Optional[str],
        container_kind: str,
        row_no: int,
        name_label: str,
) -> None:
    """按名称和类型切换容器标记，并检查嵌套顺序。

    :param active_containers: 当前从外到内尚未配对的容器栈，执行后原地更新。
    :param container_name: 标记行中的英文名称，不能为空。
    :param container_kind: 规范化后的容器类型，取 array 或 struct。
    :param row_no: 标记行在工作表中的行号，用于定位错误。
    :param name_label: 文档中英文字段名称列的表头，用于缺失提示。
    :raises InterfaceDocumentParseError: 标记缺少名称、类型冲突或关闭顺序错误。
    """
    kind_label = "数组" if container_kind == "array" else "结构体"
    if not container_name:
        raise InterfaceDocumentParseError(f"第{row_no}行{kind_label}标记缺少{name_label}")
    if active_containers and container_name == active_containers[-1]["name"]:
        if container_kind != active_containers[-1]["kind"]:
            raise InterfaceDocumentParseError(
                f"第{row_no}行容器标记「{container_name}」类型与"
                f"第{active_containers[-1]['start_row']}行不一致"
            )
        active_containers.pop()
        return
    if any(context["name"] == container_name for context in active_containers):
        raise InterfaceDocumentParseError(
            f"第{row_no}行容器标记「{container_name}」顺序不正确，"
            f"请先结束内层容器「{active_containers[-1]['name']}」"
        )
    active_containers.append({
        "name": container_name,
        "kind": container_kind,
        "start_row": row_no,
    })


def _build_field(
        *,
        field_name: str,
        field_chinese_name: Optional[str],
        required_text: Optional[str],
        length_text: Optional[str],
        enum_values: Optional[str],
        remark: Optional[str],
        array_path: Tuple[str, ...],
        source_row: int,
) -> ParsedInterfaceField:
    """将文档中的字段行转换为生成规则使用的字段定义。

    :param field_name: 字段英文名，允许英文字母和数字。
    :param field_chinese_name: 字段中文说明，可为空。
    :param required_text: 文档中的必输标记原文。
    :param length_text: 文档中的字段长度原文。
    :param enum_values: 文档中的枚举值文本，可为空。
    :param remark: 字段备注，可为空。
    :param array_path: 从外到内的所属数组与结构体名称，沿用旧快照键名。
    :param source_row: 字段在原工作表中的行号。
    :return: 含规范化规则、原文及校验错误的字段定义；非法字段名记录错误而不抛异常。
    """
    field_name_error = None
    if not _FIELD_NAME_PATTERN.fullmatch(field_name):
        field_name_error = (
            f"字段名「{field_name}」格式不合法，仅支持英文字母和数字"
        )
    length, total_length, integer_length, decimal_length = _normalize_length(length_text)
    length_error = None
    if length and total_length is None:
        length_error = f"长度「{length}」格式不合法，请检查"
    if (
            total_length is not None
            and decimal_length is not None
            and decimal_length >= total_length
    ):
        integer_length = None
        length_error = (
            f"小数长度{decimal_length}必须小于总体长度{total_length}，请检查"
        )
    required = _normalize_required(required_text)
    return {
        "field_name": field_name,
        "field_chinese_name": field_chinese_name,
        "required": required,
        "required_text": required_text,
        "length": length,
        "total_length": total_length,
        "integer_length": integer_length,
        "decimal_length": decimal_length,
        "field_name_error": field_name_error,
        "length_error": length_error,
        "enum_values": enum_values,
        "remark": remark,
        "array_path": list(array_path),
        "duplicate_in_context": False,
        "source_row": source_row,
    }


def _find_integration_headers(sheet: Worksheet) -> Tuple[int, Dict[str, int]]:
    """定位整合接口文档表头，并校验集合列允许出现的位置。"""
    header_row, row, positions = _find_ordered_header_group(
        sheet,
        INTEGRATION_HEADERS,
        integration_header_matches,
        document_name="整合接口文档",
    )
    collection_columns = []
    for index, value in enumerate(row):
        level = integration_collection_header_level(value)
        if level is not None:
            collection_columns.append((index, level))

    inner_columns = [index for index, level in collection_columns if level == 1]
    outer_columns = [index for index, level in collection_columns if level == 2]
    if len(inner_columns) != 1 or len(outer_columns) > 1:
        raise InterfaceDocumentParseError(
            "整合接口文档的所属集合/集合属性列不唯一或不存在"
        )

    inner_column = inner_columns[0]
    collection_between_names = (
        positions["字段名"] < inner_column < positions["字段描述"]
    )
    collection_after_rules = inner_column > positions["业务规则"]
    if not (collection_between_names or collection_after_rules):
        raise InterfaceDocumentParseError(
            "所属集合/集合属性只能位于字段名与字段描述之间或业务规则之后"
        )
    positions["所属集合/集合属性"] = inner_column
    if outer_columns:
        outer_column = outer_columns[0]
        if (
                outer_column <= inner_column
                or (
                    collection_between_names
                    and outer_column >= positions["字段描述"]
                )
        ):
            raise InterfaceDocumentParseError(
                "所属集合/集合属性_1必须位于所属集合/集合属性右侧且保持字段顺序"
            )
        positions["所属集合/集合属性_1"] = outer_column

    enum_columns = [
        index
        for index, value in enumerate(row)
        if enum_header_matches(value)
    ]
    if len(enum_columns) > 1:
        raise InterfaceDocumentParseError("整合接口文档存在多个枚举值列")
    if enum_columns and not (
            positions["必输/可选"]
            < enum_columns[0]
            < positions["取值范围/格式"]
    ):
        raise InterfaceDocumentParseError(
            "整合接口文档的枚举值列必须位于必输/可选与取值范围/格式之间"
        )
    return header_row, positions


def _merged_column_values(
        sheet: Worksheet,
        column_no: int,
) -> Dict[int, Optional[str]]:
    """展开指定列的纵向合并单元格值，供集合成员行继承。"""
    values: Dict[int, Optional[str]] = {}
    excel_column = column_no + 1
    for merged_range in sheet.merged_cells.ranges:
        if merged_range.min_col != excel_column or merged_range.max_col != excel_column:
            continue
        value = cell_text(sheet.cell(merged_range.min_row, excel_column).value)
        for row_no in range(merged_range.min_row, merged_range.max_row + 1):
            values[row_no] = value
    return values


def _normalize_collection_name(value: Optional[str], row_no: int) -> Optional[str]:
    """将数组名或List<数组名>统一为XML路径节点名。"""
    if not value:
        return None
    text = value.strip()
    list_match = re.fullmatch(r"List\s*<\s*([^<>]+?)\s*>", text)
    name = (list_match.group(1) if list_match else text).strip()
    if not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_.-]*", name):
        raise InterfaceDocumentParseError(f"第{row_no}行所属集合/集合属性格式不合法")
    return name


def _append_or_mark_duplicate_field(
        fields: List[ParsedInterfaceField],
        field_indexes: Dict[_FieldIdentity, int],
        field: ParsedInterfaceField,
) -> None:
    """按完整父级路径与字段名识别同一上下文中的重复字段。

    :param fields: 已解析的字段列表；首次出现的字段保留在其中。
    :param field_indexes: 字段路径到 fields 索引的映射，执行后原地更新。
    :param field: 当前待加入的字段；重复时只标记已有字段。
    """
    identity = (tuple(field["array_path"]), field["field_name"])
    existing_index = field_indexes.get(identity)
    if existing_index is not None:
        fields[existing_index]["duplicate_in_context"] = True
        return
    field_indexes[identity] = len(fields)
    fields.append(field)


def _parse_standard_sheet(
        sheet: Worksheet,
        config: _StandardSheetConfig,
) -> ParsedInterfaceDocument:
    """按标准接口文档配置解析字段和成对的数组、结构体标记。

    :param sheet: 待读取的工作表，仅解析其中首个匹配的字段表头区域。
    :param config: 当前接口样式的表头及字段列配置。
    :return: 包含接口样式、工作表信息及字段定义的解析结果。
    :raises InterfaceDocumentParseError: 表头无效或容器标记未配对、类型冲突、顺序错误。
    """
    header_row, _, positions = _find_ordered_header_group(
        sheet,
        config.headers,
        config.matches_header,
    )
    fields: List[ParsedInterfaceField] = []
    field_indexes: Dict[_FieldIdentity, int] = {}
    active_containers: List[_ContainerContext] = []

    for row_no, values in _iter_field_rows(
            sheet,
            header_row=header_row,
            positions=positions,
            field_name_header=config.field_name_header,
            chinese_name_header=config.chinese_name_header,
    ):
        field_type = (values[config.array_type_header] or "").casefold()
        if field_type in _ARRAY_TYPE_VALUES or field_type == _STRUCT_TYPE_VALUE:
            _update_container_contexts(
                active_containers,
                container_name=values[config.field_name_header],
                container_kind="array" if field_type in _ARRAY_TYPE_VALUES else "struct",
                row_no=row_no,
                name_label=config.field_name_header,
            )
            continue

        field_name = values[config.field_name_header]
        if not field_name:
            continue
        array_path = tuple(context["name"] for context in active_containers)
        field = _build_field(
            field_name=field_name,
            field_chinese_name=values[config.chinese_name_header],
            required_text=values[config.required_header],
            length_text=values[config.length_header],
            enum_values=(
                values[config.enum_header]
                if config.enum_header is not None
                else None
            ),
            remark=(
                values[config.remark_header]
                if config.remark_header is not None
                else None
            ),
            array_path=array_path,
            source_row=row_no,
        )
        _append_or_mark_duplicate_field(fields, field_indexes, field)

    if active_containers:
        active_container = active_containers[-1]
        kind_label = "数组" if active_container["kind"] == "array" else "结构体"
        raise InterfaceDocumentParseError(
            f"第{active_container['start_row']}行{kind_label}标记缺少开始或结束行"
        )
    return {
        "interface_style": config.interface_style.value,
        "sheet_name": sheet.title,
        "header_row": header_row,
        "fields": fields,
    }


def _parse_integration_sheet(sheet: Worksheet) -> ParsedInterfaceDocument:
    """解析整合接口文档，并为每个字段生成完整XML路径。"""
    header_row, positions = _find_integration_headers(sheet)
    inner_column = positions["所属集合/集合属性"]
    outer_column = positions.get("所属集合/集合属性_1")
    inner_merged_values = _merged_column_values(sheet, inner_column)
    outer_merged_values = (
        _merged_column_values(sheet, outer_column)
        if outer_column is not None
        else {}
    )

    fields: List[ParsedInterfaceField] = []
    field_indexes: Dict[_FieldIdentity, int] = {}
    for row_no, values in _iter_field_rows(
            sheet,
            header_row=header_row,
            positions=positions,
            field_name_header="字段名",
            chinese_name_header="字段描述",
    ):
        field_name = values["字段名"]
        if not field_name:
            continue

        inner_value = inner_merged_values.get(
            row_no,
            values.get("所属集合/集合属性"),
        )
        outer_value = outer_merged_values.get(
            row_no,
            values.get("所属集合/集合属性_1"),
        )
        array_name = _normalize_collection_name(inner_value, row_no)
        outer_array_name = _normalize_collection_name(outer_value, row_no)
        array_path = tuple(
            name for name in (outer_array_name, array_name) if name
        )

        enum_sources = [
            text
            for text in (values["取值范围/格式"], values["业务规则"])
            if text
        ]
        field = _build_field(
            field_name=field_name,
            field_chinese_name=values["字段描述"],
            required_text=values["必输/可选"],
            length_text=values["长度"],
            enum_values=None,
            remark="\n".join(enum_sources) or None,
            array_path=array_path,
            source_row=row_no,
        )
        _append_or_mark_duplicate_field(fields, field_indexes, field)

    return {
        "interface_style": AutoTestInterfaceStyle.INTEGRATION.value,
        "sheet_name": sheet.title,
        "header_row": header_row,
        "fields": fields,
    }


def _parse_workbook(
        content: bytes,
        interface_style: AutoTestInterfaceStyle,
) -> ParsedInterfaceDocument:
    if not content:
        raise InterfaceDocumentParseError("接口文档内容为空")
    if len(content) > MAX_DOCUMENT_SIZE:
        raise InterfaceDocumentParseError(f"接口文档大小不能超过{MAX_DOCUMENT_SIZE // 1024 // 1024}MB")
    try:
        # 整合文档需要读取纵向合并单元格范围，不能使用只读工作簿。
        workbook = load_workbook(
            io.BytesIO(content),
            read_only=interface_style != AutoTestInterfaceStyle.INTEGRATION,
            data_only=True,
        )
    except (BadZipFile, InvalidFileException, OSError, ValueError, KeyError, ParseError) as exc:
        raise InterfaceDocumentParseError("接口文档不是有效的xlsx文件") from exc

    try:
        if not workbook.sheetnames:
            raise InterfaceDocumentParseError("接口文档不包含工作表")
        sheet = workbook[workbook.sheetnames[0]]
        _validate_workbook_columns(sheet)
        if interface_style == AutoTestInterfaceStyle.INTEGRATION:
            return _parse_integration_sheet(sheet)
        return _parse_standard_sheet(
            sheet,
            _STANDARD_SHEET_CONFIGS[interface_style],
        )
    finally:
        workbook.close()


def parse_interface_document(content: bytes, interface_style: str) -> ParsedInterfaceDocument:
    """
        按标准接口样式分发解析。
    """
    try:
        style = AutoTestInterfaceStyle.normalize(interface_style)
    except ValueError as exc:
        raise InterfaceDocumentParseError(
            "接口样式仅支持esb、project或integration"
        ) from exc
    return _parse_workbook(content, style)
