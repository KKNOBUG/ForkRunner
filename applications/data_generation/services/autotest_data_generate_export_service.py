# -*- coding: utf-8 -*-

'''
    把数据库中一次数据生成任务的全部测试场景导出为excel表格，并写入本地下载目录
'''

from __future__ import annotations

import os
import tempfile
from dataclasses import dataclass
from typing import Any, Dict, List, Mapping, Optional, Sequence

from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from applications.data_generation.models.autotest_data_generate_result_model import AutoTestDataGenerateResultModel
from applications.data_generation.models.autotest_data_generate_task_model import AutoTestDataGenerateTaskModel
from applications.data_generation.schemas.autotest_data_generate_schema import normalize_generated_file_name
from applications.data_generation.services.autotest_data_generate_result_crud import AutoTestDataGenerateResultCrud
from applications.data_generation.services.autotest_data_generate_result_utils import get_result_value
from celery_scheduler.celery_task_contract import path_to_storage_key
from configure import PROJECT_CONFIG
from core.exceptions import ParameterException


@dataclass(frozen=True)
class DataGenerateExportFile:
    """数据生成导出文件定位信息，负责接收参数并把参数保存到新对象的属性中"""

    file_name: str
    file_path: str
    storage_key: str
    size: int


def _write_cell(cell: Cell, value: Optional[str]) -> None:
    """将生成数据作为文本写入Excel单元格。"""
    cell.value = "" if value is None else value
    # 避免以“=”开头的测试数据被Excel当作公式执行。
    cell.data_type = "s"


def _normalize_results(results: Sequence[Any]) -> List[Dict[str, Any]]:
    """
        在创建excel前统一检查和整理测试结果
        results：全部测试场景
        返回值：一个字典数组，键值对为字符串：任何类型，即经过检查整理的所有生成测试数据
    """
    if not results:
        raise ParameterException(message="导出数据生成结果失败, 当前任务没有可导出的测试场景")

    normalized: List[Dict[str, Any]] = []
    #已经出现过的场景名称
    seen_names = set()

    for result in results:
        #读取场景名称
        name = str(get_result_value(result, "scene_name") or "").strip()
        #！：scenario_data必须是字典结构，允许空字典
        data = get_result_value(result, "scenario_data")

        if not name:
            raise ParameterException(message="导出数据生成结果失败, 测试场景名称为空")
        if name in seen_names:
            raise ParameterException(message=f"导出数据生成结果失败, 测试场景名称重复: {name}")
        if not isinstance(data, Mapping):
            raise ParameterException(message=f"测试场景[{name}]的数据必须是字段映射")

        #当前场景名称不存在，添加进set
        seen_names.add(name)
        normalized_data: Dict[str, Any] = {}
        #提取出测试数据对应的键和值
        for raw_key, raw_value in data.items():
            path = str(raw_key)
            if path in normalized_data:
                raise ParameterException(message=f"测试场景[{name}]存在重复字段路径: {path}")
            normalized_data[path] = raw_value
        normalized.append({"scene_name": name, "scenario_data": normalized_data})
    return normalized


def build_data_generate_workbook(results: Sequence[Any]) -> Workbook:
    """
        按“字段路径为列、场景为行”创建Excel内容。
        results：全部测试场景
        Workbook：一个excel文件对象
    """
    #检查整理测试结果，得到一个键值对数组
    scenarios = _normalize_results(results)
    field_paths: List[str] = []
    seen_paths = set()
    #读取一条测试场景
    for scenario in scenarios:
        #从测试场景获取测试数据的键
        for raw_path in scenario["scenario_data"]:
            path = str(raw_path)
            #快速去重并保留原本顺序
            if path not in seen_paths:
                seen_paths.add(path)
                field_paths.append(path)
    #创建一个工作薄对象
    workbook = Workbook()
    #获取当前活动工作表
    sheet = workbook.active
    #设置sheet名
    sheet.title = "测试数据"
    #冻结窗格，方便查看
    sheet.freeze_panes = "D2"

    #表格样式
    header_fill = PatternFill("solid", fgColor="2F75B5")
    section_fill = PatternFill("solid", fgColor="D9EAF7")
    thin = Side(style="thin", color="D9E1F2")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # 数据生成仅写入BODY；保留HEAD分区列以兼容数据编辑表格结构。
    headers = ["", "Head", "Body", *field_paths]
    for column, value in enumerate(headers, start=1):
        #表示表格中对应的单元格
        cell = sheet.cell(row=1, column=column)
        _write_cell(cell, value)
        #设置字段名列的样式
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        #对齐方式/自动换行
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    #写入每一条测试场景，从第二行开始
    for row, scenario in enumerate(scenarios, start=2):
        #组装该行要写入的数据
        values = [
            scenario["scene_name"],
            "",
            "",
            *(scenario["scenario_data"].get(path) for path in field_paths),
        ]
        #写哪一列
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=row, column=column)
            _write_cell(cell, value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border

    #设置excel表格的样式
    sheet.row_dimensions[1].height = 32
    sheet.column_dimensions["A"].width = min(
        80,
        max(24, max((len(scenario["scene_name"]) for scenario in scenarios), default=12) + 2),
    )
    sheet.column_dimensions["B"].width = 12
    sheet.column_dimensions["C"].width = 12
    for column in range(4, len(field_paths) + 4):
        sheet.column_dimensions[get_column_letter(column)].width = 24
    sheet.auto_filter.ref = (
        f"A1:{get_column_letter(len(field_paths) + 3)}{len(scenarios) + 1}"
    )
    return workbook


class AutoTestDataGenerateExportService:
    """
        对外提供生成excel文件功能
    """

    def __init__(
            self,
            result_crud: Optional[AutoTestDataGenerateResultCrud] = None,
            output_root: Optional[str] = None,
    ):
        self.result_crud = result_crud or AutoTestDataGenerateResultCrud()
        self.output_root = os.path.abspath(output_root or PROJECT_CONFIG.OUTPUT_DOWNLOAD_DIR)

    @staticmethod
    def _safe_file_name(file_name: str) -> str:
        """
            验证生成的excel表格文件名
        """
        try:
            return normalize_generated_file_name(file_name)
        except ValueError as exc:
            raise ParameterException(message=str(exc)) from exc

    def _task_output_dir(self, task_code: str) -> str:
        """
            生成当前任务的专属输出目录
        """
        code = str(task_code or "").strip()
        if not code or any(char not in "abcdefghijklmnopqrstuvwxyzABCDEFGHIJKLMNOPQRSTUVWXYZ0123456789-_" for char in code):
            raise ParameterException(message="数据生成任务标识不合法")
        target = os.path.abspath(os.path.join(self.output_root, "autotest_data_generate", code))
        if os.path.commonpath((self.output_root, target)) != self.output_root:
            raise ParameterException(message="数据生成导出目录越界")
        return target

    async def export_task(
            self,
            task: AutoTestDataGenerateTaskModel,
    ) -> DataGenerateExportFile:
        """
            对外核心导出方法
        """
        results: List[AutoTestDataGenerateResultModel] = await self.result_crud.list_by_task_id(task.id)
        workbook = build_data_generate_workbook(results)

        file_name = self._safe_file_name(task.generated_file_name)
        output_dir = self._task_output_dir(task.task_code)
        os.makedirs(output_dir, exist_ok=True)
        target_path = os.path.abspath(os.path.join(output_dir, file_name))
        temp_path: Optional[str] = None
        try:
            with tempfile.NamedTemporaryFile(
                    mode="wb",
                    prefix=".data_generate_",
                    suffix=".xlsx",
                    dir=output_dir,
                    delete=False,
            ) as temp_file:
                temp_path = temp_file.name
            workbook.save(temp_path)
            os.replace(temp_path, target_path)
            temp_path = None
        finally:
            workbook.close()
            if temp_path and os.path.isfile(temp_path):
                os.unlink(temp_path)

        if self.output_root == os.path.abspath(PROJECT_CONFIG.OUTPUT_DOWNLOAD_DIR):
            storage_key = path_to_storage_key(target_path)
        else:
            storage_key = os.path.relpath(target_path, self.output_root).replace("\\", "/")
        return DataGenerateExportFile(
            file_name=file_name,
            file_path=target_path,
            storage_key=storage_key,
            size=os.path.getsize(target_path),
        )
