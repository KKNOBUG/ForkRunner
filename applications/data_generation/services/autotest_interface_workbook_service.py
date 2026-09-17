# -*- coding: utf-8 -*-

"""
    接口文档写回：将已确认的枚举值覆盖写入上传的Excel
"""

from __future__ import annotations

import os
import tempfile
from copy import copy
from dataclasses import dataclass
from typing import Any, Callable, Dict, Mapping, Optional, Sequence, Tuple
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.cell_range import CellRange
from openpyxl.worksheet.worksheet import Worksheet

from applications.data_generation.constants import (
    ENUM_EXTRACTION_INTERFACE_STYLES,
    MAX_DOCUMENT_COLUMNS,
    MAX_DOCUMENT_SIZE,
)
from applications.data_generation.services.autotest_interface_document_headers import (
    ENUM_COLUMN_HEADER,
    PROJECT_HEADER_ALIASES,
    enum_header_matches,
    integration_header_matches,
    project_header_matches,
)
from configure import PROJECT_CONFIG
from enums import AutoTestInterfaceStyle

_ENUM_ROW_HEIGHT_PER_LINE = 18


class InterfaceWorkbookError(ValueError):
    """接口文档无法安全改写。"""


@dataclass(frozen=True)
class InterfaceWorkbookOverwriteResult:
    """
        封装改写完成后的结果
    """
    file_path: str          #原始文件路径
    enum_column: int        #枚举值所在列
    inserted_column: bool   #是否插入枚举值列
    written_count: int      #实际写入枚举值的字段数量
    file_size: int          #覆盖后的文件大小


def resolve_uploaded_interface_document_path(storage_key: str) -> str:
    """
        把数据库任务表中保存的存储键转换为本地绝对路径
    """
    key = str(storage_key or "").strip().replace("\\", "/").lstrip("/")
    if not key or ".." in key.split("/"):
        raise InterfaceWorkbookError("接口文档存储路径不合法")
    root = os.path.abspath(PROJECT_CONFIG.OUTPUT_UPLOAD_DIR)
    target = os.path.abspath(os.path.join(root, *key.split("/")))
    real_root = os.path.realpath(root)
    real_target = os.path.realpath(target)
    try:
        is_inside_root = os.path.commonpath((real_root, real_target)) == real_root
    except ValueError as exc:
        raise InterfaceWorkbookError("接口文档存储路径越界") from exc
    if target == root or not is_inside_root:
        raise InterfaceWorkbookError("接口文档存储路径越界")
    return target


_HeaderMatcher = Callable[[object], bool]


def _matching_columns(
        sheet: Worksheet,
        header_row: int,
        matches_header: _HeaderMatcher,
) -> list[int]:
    """返回表头行中所有匹配的列号。"""
    return [
        column
        for column in range(1, sheet.max_column + 1)
        if matches_header(sheet.cell(header_row, column).value)
    ]


def _find_unique_column(
        sheet: Worksheet,
        header_row: int,
        matches_header: _HeaderMatcher,
        *,
        error_message: str,
) -> int:
    """查找唯一表头列，缺失或重复时返回领域错误。"""
    columns = _matching_columns(sheet, header_row, matches_header)
    if len(columns) != 1:
        raise InterfaceWorkbookError(error_message)
    return columns[0]


def _find_project_column(
        sheet: Worksheet,
        header_row: int,
        expected_header: str,
) -> int:
    """在项目接口文档表头中查找唯一目标列。"""
    names = "/".join(sorted(PROJECT_HEADER_ALIASES[expected_header]))
    return _find_unique_column(
        sheet,
        header_row,
        lambda value: project_header_matches(value, expected_header),
        error_message=f"项目接口文档表头[{names}]不唯一或不存在",
    )


def _capture_column_dimensions(sheet: Worksheet, start_column: int) -> Dict[int, Dict[str, Any]]:
    """
        保存枚举插入位置的显示属性
    """
    captured: Dict[int, Dict[str, Any]] = {}
    for column in range(start_column, sheet.max_column + 1):
        dimension = sheet.column_dimensions.get(get_column_letter(column))
        if dimension is None:
            continue
        captured[column] = {
            "width": dimension.width,
            "hidden": dimension.hidden,
            "bestFit": dimension.bestFit,
            "outlineLevel": dimension.outlineLevel,
            "collapsed": dimension.collapsed,
        }
    return captured


def _restore_shifted_column_dimensions(
        sheet: Worksheet,
        captured: Mapping[int, Mapping[str, Any]],
) -> None:
    """
        恢复移动后的列属性
    """
    for source_column in sorted(captured, reverse=True):
        target = sheet.column_dimensions[get_column_letter(source_column + 1)]
        for name, value in captured[source_column].items():
            setattr(target, name, value)


def _shift_merged_ranges(sheet: Worksheet, insert_column: int) -> list[CellRange]:
    """
        处理合并单元格
    """
    shifted: list[CellRange] = []
    for merged_range in list(sheet.merged_cells.ranges):
        cell_range = CellRange(str(merged_range))
        sheet.unmerge_cells(str(merged_range))
        if cell_range.min_col >= insert_column:
            cell_range.shift(col_shift=1)
        elif cell_range.max_col >= insert_column:
            cell_range.max_col += 1
        shifted.append(cell_range)
    return shifted


def _find_integration_value_range_column(sheet: Worksheet, header_row: int) -> int:
    """查找整合接口文档的“取值范围/格式”列。"""
    expected_header = "取值范围/格式"
    return _find_unique_column(
        sheet,
        header_row,
        lambda value: integration_header_matches(value, expected_header),
        error_message="整合接口文档表头[取值范围/格式]不唯一或不存在",
    )


def _insert_enum_column(
        sheet: Worksheet,
        insert_column: int,
        *,
        copy_source_width: bool = False,
) -> None:
    """
        将枚举值插入到新列
    """
    dimensions = _capture_column_dimensions(sheet, insert_column)
    source_width = dimensions.get(insert_column, {}).get("width")
    merged_ranges = _shift_merged_ranges(sheet, insert_column)
    sheet.insert_cols(insert_column, 1)
    _restore_shifted_column_dimensions(sheet, dimensions)
    sheet.column_dimensions[get_column_letter(insert_column)].width = (
        source_width if copy_source_width and source_width is not None else 24
    )

    # 新枚举列沿用右侧原文本列的逐行样式，避免破坏不同模板的格式。
    style_source_column = insert_column + 1
    for row in range(1, sheet.max_row + 1):
        sheet.cell(row, insert_column)._style = copy(
            sheet.cell(row, style_source_column)._style
        )
    for cell_range in merged_ranges:
        sheet.merge_cells(str(cell_range))


def _verify_written_workbook(
        file_path: str,
        *,
        header_row: int,
        enum_column: int,
        expected_values: Mapping[int, Optional[str]],
) -> None:
    """
        重新打开临时保存的excel表，验证是否成功写入枚举值
    """
    workbook = load_workbook(file_path, read_only=True, data_only=False)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        if not enum_header_matches(sheet.cell(header_row, enum_column).value):
            raise InterfaceWorkbookError("枚举值列写入校验失败")
        for row, expected in expected_values.items():
            actual = sheet.cell(row, enum_column).value
            if (None if actual is None else str(actual)) != expected:
                raise InterfaceWorkbookError(f"第{row}行枚举值写入校验失败")
    finally:
        workbook.close()


def _save_verified_workbook(
        workbook: Workbook,
        file_path: str,
        *,
        original_mode: int,
        header_row: int,
        enum_column: int,
        expected_values: Mapping[int, Optional[str]],
) -> None:
    """保存、校验并原子替换原始工作簿。"""
    temp_path: Optional[str] = None
    workbook_closed = False
    try:
        file_descriptor, temp_path = tempfile.mkstemp(
            prefix=".interface-enum-",
            suffix=".xlsx",
            dir=os.path.dirname(file_path),
        )
        os.close(file_descriptor)
        workbook.save(temp_path)
        workbook.close()
        workbook_closed = True
        _verify_written_workbook(
            temp_path,
            header_row=header_row,
            enum_column=enum_column,
            expected_values=expected_values,
        )
        os.chmod(temp_path, original_mode & 0o777)
        os.replace(temp_path, file_path)
        temp_path = None
    finally:
        if not workbook_closed:
            workbook.close()
        if temp_path and os.path.isfile(temp_path):
            try:
                os.remove(temp_path)
            except OSError:
                pass


def _validate_interface_document_snapshot(
        interface_document: Mapping[str, Any],
) -> Tuple[AutoTestInterfaceStyle, list[Any], int]:
    """校验写回所需的接口文档快照字段。"""
    try:
        interface_style = AutoTestInterfaceStyle.normalize(
            interface_document.get("interface_style")
        )
    except ValueError as exc:
        raise InterfaceWorkbookError(
            "只允许改写项目或整合接口文档"
        ) from exc
    if interface_style not in ENUM_EXTRACTION_INTERFACE_STYLES:
        raise InterfaceWorkbookError("只允许改写项目或整合接口文档")

    fields = interface_document.get("fields")
    header_row = interface_document.get("header_row")
    if not isinstance(fields, list) or not isinstance(header_row, int) or header_row < 1:
        raise InterfaceWorkbookError("接口文档快照结构不正确")
    return interface_style, fields, header_row


def _validate_interface_document_file(storage_key: str) -> Tuple[str, os.stat_result]:
    """解析上传文档路径并校验写回前的文件状态。"""
    file_path = resolve_uploaded_interface_document_path(storage_key)
    if not file_path.lower().endswith(".xlsx"):
        raise InterfaceWorkbookError("接口文档仅支持xlsx格式")
    if os.path.islink(file_path) or not os.path.isfile(file_path):
        raise InterfaceWorkbookError("接口文档不存在或不是普通文件")
    original_stat = os.stat(file_path)
    if original_stat.st_size < 1 or original_stat.st_size > MAX_DOCUMENT_SIZE:
        raise InterfaceWorkbookError("接口文档大小超出限制")
    return file_path, original_stat


def _validate_target_sheet(sheet: Worksheet, header_row: int) -> None:
    """校验目标工作表与快照表头行是否可用。"""
    if sheet.max_column > MAX_DOCUMENT_COLUMNS:
        raise InterfaceWorkbookError("接口文档列数超出限制")
    if header_row > sheet.max_row:
        raise InterfaceWorkbookError("接口文档表头行越界")


def _resolve_project_enum_column(
        sheet: Worksheet,
        header_row: int,
        existing_enum_column: Optional[int],
) -> int:
    """定位或插入项目接口文档的枚举值列。"""
    length_column = _find_project_column(sheet, header_row, "字段长度")
    remark_column = _find_project_column(sheet, header_row, "备注")
    if existing_enum_column is None:
        if length_column >= remark_column:
            raise InterfaceWorkbookError("字段长度列必须位于备注列之前")
        _insert_enum_column(sheet, remark_column)
        return remark_column
    if not length_column < existing_enum_column < remark_column:
        raise InterfaceWorkbookError("枚举值列必须位于字段长度与备注之间")
    return existing_enum_column


def _resolve_integration_enum_column(
        sheet: Worksheet,
        header_row: int,
        existing_enum_column: Optional[int],
) -> int:
    """定位或插入整合接口文档的枚举值列。"""
    value_range_column = _find_integration_value_range_column(sheet, header_row)
    if existing_enum_column is None:
        _insert_enum_column(
            sheet,
            value_range_column,
            copy_source_width=True,
        )
        return value_range_column
    if existing_enum_column >= value_range_column:
        raise InterfaceWorkbookError(
            "整合接口文档的枚举值列必须位于取值范围/格式左侧"
        )
    return existing_enum_column


def _resolve_enum_column(
        sheet: Worksheet,
        header_row: int,
        interface_style: AutoTestInterfaceStyle,
) -> Tuple[int, bool]:
    """根据接口样式定位或创建枚举值列。"""
    enum_columns = _matching_columns(sheet, header_row, enum_header_matches)
    if len(enum_columns) > 1:
        raise InterfaceWorkbookError("接口文档存在多个枚举值列")

    existing_enum_column = enum_columns[0] if enum_columns else None
    inserted_column = existing_enum_column is None
    if inserted_column and sheet.max_column >= MAX_DOCUMENT_COLUMNS:
        raise InterfaceWorkbookError("接口文档已达最大列数，无法添加枚举值列")

    if interface_style == AutoTestInterfaceStyle.PROJECT:
        enum_column = _resolve_project_enum_column(
            sheet,
            header_row,
            existing_enum_column,
        )
    else:
        enum_column = _resolve_integration_enum_column(
            sheet,
            header_row,
            existing_enum_column,
        )
    return enum_column, inserted_column


def _write_enum_values(
        sheet: Worksheet,
        fields: Sequence[Any],
        *,
        header_row: int,
        enum_column: int,
) -> Tuple[Dict[int, Optional[str]], int]:
    """将快照中的非空枚举值写入原始行并返回校验期望值。"""
    expected_values: Dict[int, Optional[str]] = {}
    written_count = 0
    for field in fields:
        if not isinstance(field, Mapping):
            raise InterfaceWorkbookError("接口字段快照结构不正确")
        source_row = field.get("source_row")
        if (
                not isinstance(source_row, int)
                or source_row <= header_row
                or source_row > sheet.max_row
        ):
            raise InterfaceWorkbookError("接口字段行号越界")

        raw_value = field.get("enum_values")
        value = str(raw_value).strip() if raw_value is not None else None
        value = value or None
        if value is None:
            continue

        cell = sheet.cell(source_row, enum_column)
        cell.value = value
        alignment = copy(cell.alignment)
        alignment.wrap_text = True
        cell.alignment = alignment
        required_height = (value.count("\n") + 1) * _ENUM_ROW_HEIGHT_PER_LINE
        row_dimension = sheet.row_dimensions[source_row]
        if required_height > (row_dimension.height or 0):
            row_dimension.height = required_height
        written_count += 1
        expected_values[source_row] = value
    return expected_values, written_count


class AutoTestInterfaceWorkbookService:
    """
        将枚举抽取结果写回任务本地接口文档并覆盖原文件。
    """

    def overwrite_interface_document(
            self,
            storage_key: str,
            interface_document: Mapping[str, Any],
    ) -> InterfaceWorkbookOverwriteResult:
        """
            将项目或整合接口的枚举抽取结果写回原文档。
            storage_key：原始接口文档在上传目录汇总的相对路径
            interface_document：枚举提取后的接口文档快照
        """
        interface_style, fields, header_row = _validate_interface_document_snapshot(
            interface_document
        )
        file_path, original_stat = _validate_interface_document_file(storage_key)

        workbook = None
        try:
            workbook = load_workbook(file_path, read_only=False, data_only=False)
            if not workbook.sheetnames:
                raise InterfaceWorkbookError("接口文档不包含工作表")
            sheet = workbook[workbook.sheetnames[0]]
            _validate_target_sheet(sheet, header_row)
            enum_column, inserted_column = _resolve_enum_column(
                sheet,
                header_row,
                interface_style,
            )

            sheet.cell(header_row, enum_column).value = ENUM_COLUMN_HEADER
            expected_values, written_count = _write_enum_values(
                sheet,
                fields,
                header_row=header_row,
                enum_column=enum_column,
            )

            workbook_to_save = workbook
            workbook = None
            _save_verified_workbook(
                workbook_to_save,
                file_path,
                original_mode=original_stat.st_mode,
                header_row=header_row,
                enum_column=enum_column,
                expected_values=expected_values,
            )
            return InterfaceWorkbookOverwriteResult(
                file_path=file_path,
                enum_column=enum_column,
                inserted_column=inserted_column,
                written_count=written_count,
                file_size=os.path.getsize(file_path),
            )
        except InterfaceWorkbookError:
            raise
        except (BadZipFile, InvalidFileException, OSError, ValueError, KeyError) as exc:
            raise InterfaceWorkbookError("接口文档改写失败") from exc
        finally:
            if workbook is not None:
                workbook.close()
